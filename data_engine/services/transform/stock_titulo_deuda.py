from dataclasses import dataclass
from datetime import date
from decimal import Decimal

import openpyxl

from data_engine.models.stock_titulo_deuda import Instrumento, Tenedor, TipoFila

FUENTE = "bcentral_emv_4"

# Índices 0-indexed (como listas/tuplas de Python): A=0 (label sector),
# B=1 (instrumento), C=2 (tenedor), D=3 (primer período de datos, 2008-I).
COLUMNA_INICIO_DATOS = 3
FILA_ANIOS = 4
FILA_TRIMESTRES = 5
FILA_INICIO_DATOS = 6
FILA_TOTAL_GENERAL = "Stock de títulos de deuda"

# Mapeo etiqueta Excel (columna B) -> Instrumento. Mantenido a mano en sync
# con el TextChoices del modelo, mismo patrón ya aceptado en otros datasets.
INSTRUMENTO_POR_ETIQUETA = {
    "Pagarés de Banco Central": Instrumento.PAGARES_BANCO_CENTRAL,
    "Certificados de depósitos a plazo": Instrumento.CERTIFICADOS_DEPOSITO_PLAZO,
    "Efectos de comercio": Instrumento.EFECTOS_COMERCIO,
    "Bonos de Banco Central": Instrumento.BONOS_BANCO_CENTRAL,
    "Bonos de Bancos y cooperativas": Instrumento.BONOS_BANCOS_COOPERATIVAS,
    "Bonos de Otros intermediarios financieros": Instrumento.BONOS_OTROS_INTERMEDIARIOS_FINANCIEROS,
    "Bonos de Empresas no financieras": Instrumento.BONOS_EMPRESAS_NO_FINANCIERAS,
    "Bonos de Gobierno general": Instrumento.BONOS_GOBIERNO_GENERAL,
    "Bonos de No residentes (Mercado local)": Instrumento.BONOS_NO_RESIDENTES,
}

# Mapeo etiqueta Excel (columna C, ya con .strip() aplicado) -> Tenedor.
TENEDOR_POR_ETIQUETA = {
    "Banco Central": Tenedor.BANCO_CENTRAL,
    "Bancos y cooperativas": Tenedor.BANCOS_COOPERATIVAS,
    "Fondos mutuos y de inversión": Tenedor.FONDOS_MUTUOS_INVERSION,
    "Otros intermediarios financieros": Tenedor.OTROS_INTERMEDIARIOS_FINANCIEROS,
    "Fondos de pensiones": Tenedor.FONDOS_PENSIONES,
    "Compañias de seguros": Tenedor.COMPANIAS_SEGUROS,
    "Gobierno general": Tenedor.GOBIERNO_GENERAL,
    "Otros sectores residentes": Tenedor.OTROS_SECTORES_RESIDENTES,
    "Inversionistas extranjeros": Tenedor.INVERSIONISTAS_EXTRANJEROS,
}

PREFIJO_SUBCONJUNTO = "de lo cual"


@dataclass
class RegistroStockTituloDeuda:
    instrumento: str
    tenedor: str | None
    tipo_fila: str
    anio: int
    trimestre: int
    fecha_corte: date
    monto_miles_millones_clp: Decimal


def _fecha_fin_trimestre(anio: int, trimestre: int) -> date:
    """Último día del trimestre: I->31-mar, II->30-jun, III->30-sep, IV->31-dic."""
    ultimo_dia_por_trimestre = {1: (3, 31), 2: (6, 30), 3: (9, 30), 4: (12, 31)}
    mes, dia = ultimo_dia_por_trimestre[trimestre]
    return date(anio, mes, dia)


def _leer_periodos(ws) -> list[tuple[int, int, date]]:
    """Lee las filas de año/trimestre y devuelve (anio, trimestre, fecha_corte) por columna de datos."""
    fila_anios = [c.value for c in ws[FILA_ANIOS]]
    fila_trimestres = [c.value for c in ws[FILA_TRIMESTRES]]

    periodos = []
    anio_actual = None
    for col_idx in range(COLUMNA_INICIO_DATOS, len(fila_anios)):
        if fila_anios[col_idx] is not None:
            anio_actual = int(fila_anios[col_idx])
        trimestre_romano = fila_trimestres[col_idx]
        if trimestre_romano is None:
            continue
        trimestre = {"I": 1, "II": 2, "III": 3, "IV": 4}[trimestre_romano]
        periodos.append((anio_actual, trimestre, _fecha_fin_trimestre(anio_actual, trimestre)))
    return periodos


def _parsear_monto(valor) -> Decimal:
    if valor is None:
        return Decimal("0")
    return Decimal(str(valor))


def _validar_contra_total(ws, periodos: list[tuple[int, int, date]],
                           registros: list["RegistroStockTituloDeuda"],
                           tolerancia: Decimal = Decimal("0.01")) -> None:
    """Compara la suma de todas las filas TOTAL_INSTRUMENTO por período
    contra la fila 'Stock de títulos de deuda' (gran total) de la fuente."""
    fila_total = next(
        fila for fila in ws.iter_rows(min_row=FILA_INICIO_DATOS, max_row=ws.max_row)
        if fila[1].value == FILA_TOTAL_GENERAL
    )
    totales_fuente = fila_total[COLUMNA_INICIO_DATOS:COLUMNA_INICIO_DATOS + len(periodos)]

    suma_por_periodo: dict[tuple[int, int], Decimal] = {}
    for r in registros:
        if r.tipo_fila == TipoFila.TOTAL_INSTRUMENTO:
            clave = (r.anio, r.trimestre)
            suma_por_periodo[clave] = suma_por_periodo.get(clave, Decimal("0")) + r.monto_miles_millones_clp

    for (anio, trimestre, _fecha), celda_total in zip(periodos, totales_fuente):
        suma_calculada = suma_por_periodo[(anio, trimestre)]
        total_fuente = _parsear_monto(celda_total.value)
        if abs(suma_calculada - total_fuente) > tolerancia:
            raise ValueError(
                f"Total no calza para {anio}-T{trimestre}: "
                f"suma_instrumentos={suma_calculada}, total_fuente={total_fuente}"
            )


def transformar(path_excel: str) -> list[RegistroStockTituloDeuda]:
    wb = openpyxl.load_workbook(path_excel, data_only=True)
    ws = wb["EMV_4"]

    periodos = _leer_periodos(ws)
    num_periodos = len(periodos)

    registros: list[RegistroStockTituloDeuda] = []
    instrumento_actual: str | None = None

    for fila in ws.iter_rows(min_row=FILA_INICIO_DATOS, max_row=ws.max_row):
        etiqueta_instrumento = fila[1].value  # columna B
        etiqueta_tenedor_raw = fila[2].value  # columna C

        if etiqueta_instrumento == FILA_TOTAL_GENERAL:
            continue  # fila de gran total: no se persiste, solo se usa para validar

        if etiqueta_instrumento is not None:
            # Fila de instrumento: fija el instrumento actual y persiste su total
            instrumento_actual = INSTRUMENTO_POR_ETIQUETA[etiqueta_instrumento]
            valores = fila[COLUMNA_INICIO_DATOS:COLUMNA_INICIO_DATOS + num_periodos]
            for (anio, trimestre, fecha_corte), celda in zip(periodos, valores):
                registros.append(RegistroStockTituloDeuda(
                    instrumento=instrumento_actual,
                    tenedor=None,
                    tipo_fila=TipoFila.TOTAL_INSTRUMENTO,
                    anio=anio, trimestre=trimestre, fecha_corte=fecha_corte,
                    monto_miles_millones_clp=_parsear_monto(celda.value),
                ))
            continue

        if etiqueta_tenedor_raw is not None:
            etiqueta_tenedor = etiqueta_tenedor_raw.strip()
            es_subconjunto = etiqueta_tenedor.lower().startswith(PREFIJO_SUBCONJUNTO)
            if es_subconjunto:
                # "   de lo cual: Comprado en el ML" -> se guarda bajo el tenedor de la fila de arriba
                # (Inversionistas extranjeros), pero marcado como SUBCONJUNTO para no sumarlo con TENEDOR
                tenedor = Tenedor.INVERSIONISTAS_EXTRANJEROS
                tipo_fila = TipoFila.SUBCONJUNTO
            else:
                tenedor = TENEDOR_POR_ETIQUETA[etiqueta_tenedor]
                tipo_fila = TipoFila.TENEDOR

            valores = fila[COLUMNA_INICIO_DATOS:COLUMNA_INICIO_DATOS + num_periodos]
            for (anio, trimestre, fecha_corte), celda in zip(periodos, valores):
                registros.append(RegistroStockTituloDeuda(
                    instrumento=instrumento_actual,
                    tenedor=tenedor,
                    tipo_fila=tipo_fila,
                    anio=anio, trimestre=trimestre, fecha_corte=fecha_corte,
                    monto_miles_millones_clp=_parsear_monto(celda.value),
                ))

    _validar_contra_total(ws, periodos, registros)
    return registros